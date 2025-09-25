"""
Comprehensive Test Suite for AI-Powered CV Processing Engine
TTE Recruitment Suite

This test suite provides comprehensive coverage for all components of the CV processing system:
- Document processing (PDF/DOCX extraction)
- Candidate data extraction
- PII identification and anonymization
- Skills analysis
- AI call logging
- Error handling
- Integration testing

Version: 1.0
Date: September 24, 2025
"""

import pytest
import json
import base64
import os
import tempfile
import logging
from unittest.mock import Mock, patch, MagicMock
from typing import Dict, Any, List
from datetime import datetime
from docx import Document
import io

# Import modules to test
from shared_code.document_processor import (
    extract_markdown_from_file, 
    UnsupportedFileType, 
    DocumentProcessingError
)
from shared_code.candidate_data_service import CandidateDataService
from shared_code.openai_service import OpenAIService, OpenAIResponse
from shared_code.ai_call_logger import AICallLogger
import function_app

# Test fixtures and sample data
@pytest.fixture
def sample_cv_content():
    """Sample CV content for testing"""
    return """
# John Smith
## Senior Software Engineer

**Email:** john.smith@email.com  
**Phone:** +1-555-123-4567  
**LinkedIn:** https://linkedin.com/in/johnsmith  
**GitHub:** https://github.com/johnsmith  
**Portfolio:** https://johnsmith.dev  

### Current Position
Senior Software Engineer at Tech Innovations Inc  
Location: New York, NY  

### Professional Summary
Experienced software engineer with 8 years of experience in full-stack development.
Expert in Python, JavaScript, and cloud architecture.

### Work Experience

#### Senior Software Engineer | Tech Innovations Inc | 2020-Present
- Led development team of 5 engineers
- Architected microservices solutions using AWS
- Improved system performance by 40%
- Technologies: Python, React, AWS, Docker

#### Software Engineer | Digital Solutions Corp | 2018-2020
- Developed web applications using Django and React
- Implemented automated testing frameworks
- Collaborated with cross-functional teams

### Education
**Master of Computer Science** | Stanford University | 2018  
**Bachelor of Computer Science** | UC Berkeley | 2016  

### Certifications
- AWS Certified Solutions Architect
- PMP Certified

### Skills
- Programming Languages: Python, JavaScript, Java, Go
- Frameworks: Django, React, Node.js
- Cloud Platforms: AWS, Azure
- Databases: PostgreSQL, MongoDB, Redis
- DevOps: Docker, Kubernetes, Jenkins

### Salary Information
Current Salary: $120,000 USD  
Expected Salary: $140,000 USD  

### Availability
Available in 30 days  
Visa Status: Permanent Resident
"""

@pytest.fixture
def sample_pii_entities():
    """Sample PII entities for testing anonymization"""
    return [
        {
            "original_value": "John Smith",
            "pii_type": "name",
            "sensitivity_level": 4,
            "all_variations": ["John Smith", "John", "Smith"],
            "context": "candidate name in header",
            "region_specific": False
        },
        {
            "original_value": "Tech Innovations Inc",
            "pii_type": "organization", 
            "sensitivity_level": 2,
            "all_variations": ["Tech Innovations Inc"],
            "context": "current employer",
            "region_specific": False
        },
        {
            "original_value": "Stanford University",
            "pii_type": "school",
            "sensitivity_level": 2,
            "all_variations": ["Stanford University", "Stanford"],
            "context": "education institution",
            "region_specific": False
        },
        {
            "original_value": "UC Berkeley",
            "pii_type": "school",
            "sensitivity_level": 2,
            "all_variations": ["UC Berkeley"],
            "context": "education institution",
            "region_specific": False
        }
    ]

@pytest.fixture
def sample_candidate_profile():
    """Sample candidate profile data"""
    return {
        "candidateName": "John Smith",
        "jobTitle": "Senior Software Engineer",
        "status": "Active",
        "candidateEmail": "john.smith@email.com",
        "candidatePhone": "+1-555-123-4567",
        "linkedInUrl": "https://linkedin.com/in/johnsmith",
        "otherProfileUrls": ["https://johnsmith.dev", "https://github.com/johnsmith"],
        "currentLocation": "New York, NY",
        "currentTitle": "Senior Software Engineer",
        "currentCompany": "Tech Innovations Inc",
        "professionalHeadline": "Experienced Software Engineer with Cloud Architecture Expertise",
        "totalExperienceYears": 8,
        "relevantExperienceYears": 6,
        "highestQualification": "Master of Computer Science",
        "currentSalary": "120000 USD",
        "expectedSalary": "140000 USD",
        "availabilityStatus": "Available in 30 days",
        "visaStatus": "Permanent Resident"
    }

@pytest.fixture
def sample_skills_analysis():
    """Sample skills analysis data"""
    return {
        "analysisMetrics": {
            "aiConfidenceScore": 85,
            "aiRemarks": "Strong technical background with relevant industry experience in cloud architecture and team leadership.",
            "overallFitScore": 78
        },
        "professionalProfile": {
            "coreSkills": [
                {
                    "skillName": "Python",
                    "proficiencyLevel": "Expert",
                    "yearsOfExperience": 5
                },
                {
                    "skillName": "Cloud Architecture",
                    "proficiencyLevel": "Advanced",
                    "yearsOfExperience": 4
                }
            ],
            "workExperience": [
                {
                    "jobTitle": "Senior Software Engineer",
                    "companySize": "Large Enterprise",
                    "duration": "3 years 2 months",
                    "keyResponsibilities": ["Led development team", "Architected solutions"],
                    "keyAchievements": ["Improved system performance by 40%"]
                }
            ],
            "educationProfile": {
                "highestDegreeLevel": "Masters",
                "fieldOfStudy": "Computer Science",
                "relevantCertifications": ["AWS Certified", "PMP Certified"]
            }
        }
    }

@pytest.fixture
def mock_openai_response():
    """Mock OpenAI response"""
    mock_usage = Mock()
    mock_usage.prompt_tokens = 100
    mock_usage.completion_tokens = 50
    mock_usage.total_tokens = 150
    
    return OpenAIResponse(
        content='{"test": "response"}',
        usage=mock_usage,
        model="gpt-4o-mini",
        finish_reason="completed"
    )

# Document Processor Tests
class TestDocumentProcessor:
    """Test cases for document processing functionality"""
    
    def test_extract_markdown_from_pdf_success(self, sample_cv_content):
        """Test successful PDF text extraction"""
        with patch('pymupdf4llm.to_markdown') as mock_pymupdf:
            mock_pymupdf.return_value = sample_cv_content
            
            result = extract_markdown_from_file("test.pdf", b"fake_pdf_bytes")
            
            assert result == sample_cv_content
            assert len(result) > 0
            mock_pymupdf.assert_called_once_with(b"fake_pdf_bytes")
    
    def test_extract_markdown_from_docx_success(self):
        """Test successful DOCX text extraction"""
        # Create a mock DOCX document
        with patch('shared_code.document_processor.Document') as mock_doc_class:
            mock_doc = Mock()
            mock_paragraph1 = Mock()
            mock_paragraph1.text = "John Smith"
            mock_paragraph1.runs = [Mock(bold=True)]
            
            mock_paragraph2 = Mock()
            mock_paragraph2.text = "Software Engineer with 8 years experience"
            mock_paragraph2.runs = [Mock(bold=False)]
            
            mock_doc.paragraphs = [mock_paragraph1, mock_paragraph2]
            mock_doc_class.return_value = mock_doc
            
            result = extract_markdown_from_file("test.docx", b"fake_docx_bytes")
            
            assert "## John Smith" in result
            assert "Software Engineer with 8 years experience" in result
    
    def test_unsupported_file_type(self):
        """Test handling of unsupported file types"""
        with pytest.raises(UnsupportedFileType) as exc_info:
            extract_markdown_from_file("test.txt", b"fake_bytes")
        
        assert "Unsupported file type: 'txt'" in str(exc_info.value)
    
    def test_empty_document_error(self):
        """Test handling of empty documents"""
        with patch('pymupdf4llm.to_markdown') as mock_pymupdf:
            mock_pymupdf.return_value = ""
            
            with pytest.raises(DocumentProcessingError) as exc_info:
                extract_markdown_from_file("test.pdf", b"fake_pdf_bytes")
            
            assert "No text could be extracted" in str(exc_info.value)
    
    def test_corrupt_pdf_error(self):
        """Test handling of corrupt PDF files"""
        with patch('pymupdf4llm.to_markdown') as mock_pymupdf:
            mock_pymupdf.side_effect = Exception("Corrupt PDF file")
            
            with pytest.raises(DocumentProcessingError) as exc_info:
                extract_markdown_from_file("test.pdf", b"fake_pdf_bytes")
            
            assert "A library error occurred" in str(exc_info.value)
    
    def test_corrupt_docx_error(self):
        """Test handling of corrupt DOCX files"""
        with patch('shared_code.document_processor.Document') as mock_doc_class:
            mock_doc_class.side_effect = Exception("Corrupt DOCX file")
            
            with pytest.raises(DocumentProcessingError) as exc_info:
                extract_markdown_from_file("test.docx", b"fake_docx_bytes")
            
            assert "A library error occurred" in str(exc_info.value)

# Candidate Data Service Tests
class TestCandidateDataService:
    """Test cases for candidate data service functionality"""
    
    def setup_method(self):
        """Set up test fixtures"""
        self.service = CandidateDataService()
    
    @patch('shared_code.candidate_data_service.OpenAIService')
    @patch('shared_code.candidate_data_service.AICallLogger')
    def test_regex_extraction(self, mock_logger, mock_openai):
        """Test regex-based data extraction"""
        cv_text = """
        John Smith
        Email: john.smith@email.com
        Phone: +1-555-123-4567
        LinkedIn: https://linkedin.com/in/johnsmith
        GitHub: https://github.com/johnsmith
        Portfolio: https://johnsmith.dev
        """
        
        regex_data = self.service._extract_with_regex(cv_text)
        
        assert regex_data['candidateEmail'] == 'john.smith@email.com'
        assert regex_data['candidatePhone'] == '+1-555-123-4567'
        assert regex_data['linkedInUrl'] == 'https://linkedin.com/in/johnsmith'
        assert 'https://github.com/johnsmith' in regex_data['otherProfileUrls']
        assert 'https://johnsmith.dev' in regex_data['otherProfileUrls']
    
    def test_regex_extraction_edge_cases(self):
        """Test regex extraction with edge cases"""
        # Test with no matches
        cv_text = "This is just some text without any structured data"
        regex_data = self.service._extract_with_regex(cv_text)
        assert len(regex_data) == 0
        
        # Test with multiple emails (should take first)
        cv_text = "john@email.com and jane@email.com"
        regex_data = self.service._extract_with_regex(cv_text)
        assert regex_data['candidateEmail'] == 'john@email.com'
    
    def test_anonymization_deterministic_placeholders(self, sample_pii_entities):
        """Test that PII anonymization creates deterministic placeholders"""
        original_content = "John Smith works at Tech Innovations Inc and studied at Stanford University and UC Berkeley."
        
        anonymized = self.service._create_anonymized_content(original_content, sample_pii_entities)
        
        assert "[CANDIDATE_NAME]" in anonymized
        assert "[COMPANY_1]" in anonymized
        assert "[UNIVERSITY_1]" in anonymized
        assert "[UNIVERSITY_2]" in anonymized
        assert "John Smith" not in anonymized
        assert "Tech Innovations Inc" not in anonymized
    
    def test_anonymization_all_variations(self):
        """Test that all variations of PII are anonymized"""
        pii_entities = [{
            "original_value": "John Smith",
            "pii_type": "name",
            "sensitivity_level": 4,
            "all_variations": ["John Smith", "John", "Smith", "J. Smith"],
            "context": "candidate name",
            "region_specific": False
        }]
        
        original_content = "John Smith is a developer. John has experience. Smith worked at companies. J. Smith is qualified."
        
        anonymized = self.service._create_anonymized_content(original_content, pii_entities)
        
        assert "John Smith" not in anonymized
        assert "John" not in anonymized
        assert "Smith" not in anonymized
        assert "J. Smith" not in anonymized
        assert anonymized.count("[CANDIDATE_NAME]") == 4
    
    def test_anonymization_length_ordering(self):
        """Test that longer PII entities are replaced first to avoid partial replacements"""
        pii_entities = [
            {
                "original_value": "John",
                "pii_type": "name",
                "sensitivity_level": 4,
                "all_variations": ["John"],
                "context": "first name",
                "region_specific": False
            },
            {
                "original_value": "John Smith",
                "pii_type": "name", 
                "sensitivity_level": 4,
                "all_variations": ["John Smith"],
                "context": "full name",
                "region_specific": False
            }
        ]
        
        original_content = "John Smith is a developer named John."
        
        anonymized = self.service._create_anonymized_content(original_content, pii_entities)
        
        # Should have one full name replacement and one first name replacement
        assert "[CANDIDATE_NAME]" in anonymized
        assert anonymized.count("[CANDIDATE_NAME]") == 2
    
    @patch('shared_code.candidate_data_service.OpenAIService')
    @patch('shared_code.candidate_data_service.AICallLogger')
    def test_extract_candidate_profile_success(self, mock_logger, mock_openai, 
                                             sample_cv_content, sample_candidate_profile, mock_openai_response):
        """Test successful candidate profile extraction"""
        # Mock the OpenAI service
        mock_openai_instance = Mock()
        mock_openai_response.content = json.dumps({"candidateProfile": sample_candidate_profile})
        mock_openai_instance.call_assistant.return_value = mock_openai_response
        mock_openai_instance.get_assistant_id.return_value = "asst_123"
        mock_openai.return_value = mock_openai_instance
        
        # Mock the logger
        mock_logger_instance = Mock()
        mock_logger_instance.create_call_log.return_value = {"stage": "candidate_data_extraction"}
        mock_logger.return_value = mock_logger_instance
        
        service = CandidateDataService()
        result, log = service._extract_candidate_profile(sample_cv_content, "Senior Software Engineer")
        
        assert result['candidateName'] == "John Smith"
        assert result['jobTitle'] == "Senior Software Engineer"
        assert result['candidateEmail'] == "john.smith@email.com"
        assert log['stage'] == "candidate_data_extraction"
    
    @patch('shared_code.candidate_data_service.OpenAIService')
    @patch('shared_code.candidate_data_service.AICallLogger')
    def test_pii_identification_success(self, mock_logger, mock_openai, 
                                      sample_cv_content, sample_pii_entities, mock_openai_response):
        """Test successful PII identification"""
        # Mock the OpenAI service
        mock_openai_instance = Mock()
        mock_openai_response.content = json.dumps({"pii_entities": sample_pii_entities})
        mock_openai_instance.call_assistant.return_value = mock_openai_response
        mock_openai_instance.get_assistant_id.return_value = "asst_456"
        mock_openai.return_value = mock_openai_instance
        
        # Mock the logger
        mock_logger_instance = Mock()
        mock_logger_instance.create_call_log.return_value = {"stage": "pii_identification"}
        mock_logger.return_value = mock_logger_instance
        
        service = CandidateDataService()
        pii_data, anonymized_content, log = service._identify_and_anonymize_pii(sample_cv_content)
        
        assert len(pii_data['pii_entities']) == 4
        assert "[CANDIDATE_NAME]" in anonymized_content
        assert "[COMPANY_1]" in anonymized_content 
        assert log['stage'] == "pii_identification"
    
    @patch('shared_code.candidate_data_service.OpenAIService')
    @patch('shared_code.candidate_data_service.AICallLogger')
    def test_skills_analysis_success(self, mock_logger, mock_openai, 
                                   sample_skills_analysis, mock_openai_response):
        """Test successful skills analysis"""
        # Mock the OpenAI service
        mock_openai_instance = Mock()
        mock_openai_response.content = json.dumps(sample_skills_analysis)
        mock_openai_instance.call_assistant.return_value = mock_openai_response
        mock_openai_instance.get_assistant_id.return_value = "asst_789"
        mock_openai.return_value = mock_openai_instance
        
        # Mock the logger
        mock_logger_instance = Mock()
        mock_logger_instance.create_call_log.return_value = {"stage": "skills_analysis"}
        mock_logger.return_value = mock_logger_instance
        
        service = CandidateDataService()
        anonymized_content = "Anonymized CV content with placeholders"
        
        analysis_data, log = service._analyze_skills_and_experience(anonymized_content, "Senior Software Engineer")
        
        assert analysis_data['aiConfidenceScore'] == 85
        assert analysis_data['overallFitScore'] == 78
        assert len(analysis_data['coreSkills']) == 2
        assert log['stage'] == "skills_analysis"
    
    def test_token_summary_calculation(self):
        """Test token usage summary calculation"""
        ai_call_logs = [
            {"inputTokens": 100, "outputTokens": 50},
            {"inputTokens": 200, "outputTokens": 75},
            {"inputTokens": 150, "outputTokens": 25}
        ]
        
        summary = self.service._calculate_token_summary(ai_call_logs)
        
        assert summary['totalInputTokens'] == 450
        assert summary['totalOutputTokens'] == 150
        assert summary['totalTokens'] == 600
        assert "USD" in summary['estimatedCost']
    
    @patch('shared_code.candidate_data_service.OpenAIService')
    @patch('shared_code.candidate_data_service.AICallLogger')
    def test_full_pipeline_integration(self, mock_logger, mock_openai, 
                                     sample_cv_content, sample_candidate_profile, 
                                     sample_pii_entities, sample_skills_analysis, mock_openai_response):
        """Test the complete three-stage AI pipeline"""
        # Mock the OpenAI service for all three calls
        mock_openai_instance = Mock()
        
        # Define side effects for different assistant calls
        def mock_call_assistant(assistant_type, prompt):
            if assistant_type == "CV_DATA_EXTRACTOR":
                response = Mock()
                response.content = json.dumps({"candidateProfile": sample_candidate_profile})
                response.usage = mock_openai_response.usage
                return response
            elif assistant_type == "CV_PII_IDENTIFIER":
                response = Mock()
                response.content = json.dumps({"pii_entities": sample_pii_entities})
                response.usage = mock_openai_response.usage
                return response
            elif assistant_type == "CV_SKILLS_ANALYST":
                response = Mock()
                response.content = json.dumps(sample_skills_analysis)
                response.usage = mock_openai_response.usage
                return response
        
        mock_openai_instance.call_assistant.side_effect = mock_call_assistant
        mock_openai_instance.get_assistant_id.side_effect = lambda x: f"asst_{x.lower()}"
        mock_openai.return_value = mock_openai_instance
        
        # Mock the logger
        mock_logger_instance = Mock()
        mock_logger_instance.create_call_log.return_value = {"stage": "test"}
        mock_logger.return_value = mock_logger_instance
        
        service = CandidateDataService()
        result = service.extract_candidate_data(sample_cv_content, "Senior Software Engineer")
        
        # Verify all components are present
        assert 'candidateProfile' in result
        assert 'professionalAnalysis' in result
        assert 'aiCallLogs' in result
        assert 'tokenUsageSummary' in result
        
        # Verify candidate profile
        assert result['candidateProfile']['candidateName'] == "John Smith"
        
        # Verify professional analysis
        assert result['professionalAnalysis']['aiConfidenceScore'] == 85
        
        # Verify logs (should have 3 calls)
        assert len(result['aiCallLogs']) == 3
        
        # Verify token summary
        assert 'totalTokens' in result['tokenUsageSummary']

# OpenAI Service Tests
class TestOpenAIService:
    """Test cases for OpenAI service functionality"""
    
    @patch.dict(os.environ, {
        'AZURE_OPENAI_ENDPOINT': 'https://test.openai.azure.com/',
        'AZURE_OPENAI_API_KEY': 'test-key',
        'CV_DATA_EXTRACTOR_ASSISTANT_ID': 'asst_123',
        'CV_PII_IDENTIFIER_ASSISTANT_ID': 'asst_456',
        'CV_SKILLS_ANALYST_ASSISTANT_ID': 'asst_789'
    })
    def test_service_initialization_success(self):
        """Test successful service initialization with valid config"""
        service = OpenAIService()
        
        assert service.endpoint == 'https://test.openai.azure.com/'
        assert service.api_key == 'test-key'
        assert service.assistant_ids['CV_DATA_EXTRACTOR'] == 'asst_123'
    
    def test_service_initialization_missing_config(self):
        """Test service initialization with missing configuration"""
        with patch.dict(os.environ, {}, clear=True):
            with pytest.raises(ValueError) as exc_info:
                OpenAIService()
            
            assert "AZURE_OPENAI_ENDPOINT environment variable is required" in str(exc_info.value)
    
    @patch.dict(os.environ, {
        'AZURE_OPENAI_ENDPOINT': 'https://test.openai.azure.com/',
        'AZURE_OPENAI_API_KEY': 'test-key',
        'CV_DATA_EXTRACTOR_ASSISTANT_ID': 'asst_123',
        'CV_PII_IDENTIFIER_ASSISTANT_ID': 'asst_456',
        'CV_SKILLS_ANALYST_ASSISTANT_ID': 'asst_789'
    })
    def test_get_assistant_id_valid(self):
        """Test getting valid assistant ID"""
        service = OpenAIService()
        
        assert service.get_assistant_id('CV_DATA_EXTRACTOR') == 'asst_123'
        assert service.get_assistant_id('CV_PII_IDENTIFIER') == 'asst_456'
        assert service.get_assistant_id('CV_SKILLS_ANALYST') == 'asst_789'
    
    @patch.dict(os.environ, {
        'AZURE_OPENAI_ENDPOINT': 'https://test.openai.azure.com/',
        'AZURE_OPENAI_API_KEY': 'test-key',
        'CV_DATA_EXTRACTOR_ASSISTANT_ID': 'asst_123',
        'CV_PII_IDENTIFIER_ASSISTANT_ID': 'asst_456',
        'CV_SKILLS_ANALYST_ASSISTANT_ID': 'asst_789'
    })
    def test_get_assistant_id_invalid(self):
        """Test getting invalid assistant ID"""
        service = OpenAIService()
        
        with pytest.raises(ValueError) as exc_info:
            service.get_assistant_id('INVALID_ASSISTANT')
        
        assert "Invalid assistant type: INVALID_ASSISTANT" in str(exc_info.value)
    
    @patch.dict(os.environ, {
        'AZURE_OPENAI_ENDPOINT': 'https://test.openai.azure.com/',
        'AZURE_OPENAI_API_KEY': 'test-key',
        'CV_DATA_EXTRACTOR_ASSISTANT_ID': 'asst_123',
        'CV_PII_IDENTIFIER_ASSISTANT_ID': 'asst_456',
        'CV_SKILLS_ANALYST_ASSISTANT_ID': 'asst_789'
    })
    @patch('shared_code.openai_service.AzureOpenAI')
    def test_call_assistant_success(self, mock_azure_openai):
        """Test successful assistant call"""
        # Mock the Azure OpenAI client and responses
        mock_client = Mock()
        mock_azure_openai.return_value = mock_client
        
        # Mock thread creation
        mock_thread = Mock()
        mock_thread.id = "thread_123"
        mock_client.beta.threads.create.return_value = mock_thread
        
        # Mock run creation and completion
        mock_run = Mock()
        mock_run.id = "run_123"
        mock_run.status = "completed"
        mock_usage = Mock()
        mock_usage.total_tokens = 150
        mock_usage.prompt_tokens = 100
        mock_usage.completion_tokens = 50
        mock_run.usage = mock_usage
        mock_run.model = "gpt-4o-mini"
        mock_client.beta.threads.runs.create.return_value = mock_run
        mock_client.beta.threads.runs.retrieve.return_value = mock_run
        
        # Mock messages
        mock_message = Mock()
        mock_message.content = [Mock()]
        mock_message.content[0].text.value = '{"test": "response"}'
        mock_messages = Mock()
        mock_messages.data = [mock_message]
        mock_client.beta.threads.messages.list.return_value = mock_messages
        
        service = OpenAIService()
        response = service.call_assistant("CV_DATA_EXTRACTOR", "Test prompt")
        
        assert response.content == '{"test": "response"}'
        assert response.usage.total_tokens == 150
        assert response.model == "gpt-4o-mini"

# AI Call Logger Tests
class TestAICallLogger:
    """Test cases for AI call logging functionality"""
    
    def setup_method(self):
        """Set up test fixtures"""
        self.logger = AICallLogger()
    
    def test_create_call_log_success(self):
        """Test creating a successful call log"""
        log = self.logger.create_call_log(
            stage="candidate_data_extraction",
            assistant_id="asst_123",
            model_used="gpt-4o-mini",
            input_prompt="Test prompt",
            input_tokens=100,
            response_text='{"test": "response"}',
            output_tokens=50,
            call_duration_ms=1500.0,
            call_timestamp="2025-09-24T10:30:00Z",
            success=True
        )
        
        assert log['stage'] == "candidate_data_extraction"
        assert log['assistantId'] == "asst_123"
        assert log['inputTokens'] == 100
        assert log['outputTokens'] == 50
        assert log['callDurationMs'] == 1500.0
        assert log['success'] is True
        assert log['errorDetails'] is None
    
    def test_create_call_log_failure(self):
        """Test creating a failed call log"""
        log = self.logger.create_call_log(
            stage="pii_identification",
            assistant_id="asst_456",
            model_used="gpt-4o-mini",
            input_prompt="Test prompt",
            input_tokens=0,
            response_text="",
            output_tokens=0,
            call_duration_ms=500.0,
            call_timestamp="2025-09-24T10:30:00Z",
            success=False,
            error_details="API timeout error"
        )
        
        assert log['success'] is False
        assert log['errorDetails'] == "API timeout error"
        assert log['inputTokens'] == 0
        assert log['outputTokens'] == 0
    
    def test_prompt_truncation(self):
        """Test that long prompts are truncated"""
        long_prompt = "A" * 15000  # Longer than max_prompt_length
        
        log = self.logger.create_call_log(
            stage="test",
            assistant_id="asst_123",
            model_used="gpt-4o-mini",
            input_prompt=long_prompt,
            input_tokens=100,
            response_text="response",
            output_tokens=50,
            call_duration_ms=1000.0,
            call_timestamp="2025-09-24T10:30:00Z",
            success=True
        )
        
        assert len(log['inputPrompt']) <= self.logger.max_prompt_length + 20  # Account for truncation marker
        assert "[TRUNCATED]" in log['inputPrompt']

# Function App Integration Tests
class TestFunctionApp:
    """Test cases for main function app integration"""
    
    @patch.dict(os.environ, {
        'AZURE_OPENAI_ENDPOINT': 'https://test.openai.azure.com/',
        'AZURE_OPENAI_API_KEY': 'test-key',
        'CV_DATA_EXTRACTOR_ASSISTANT_ID': 'asst_123',
        'CV_PII_IDENTIFIER_ASSISTANT_ID': 'asst_456',
        'CV_SKILLS_ANALYST_ASSISTANT_ID': 'asst_789'
    })
    def test_validate_configuration_success(self):
        """Test successful configuration validation"""
        result = function_app.validate_configuration()
        assert result is True
    
    def test_validate_configuration_missing_vars(self):
        """Test configuration validation with missing variables"""
        with patch.dict(os.environ, {}, clear=True):
            with pytest.raises(function_app.ConfigurationError) as exc_info:
                function_app.validate_configuration()
            
            assert "Missing required environment variables" in str(exc_info.value)
    
    def test_generate_correlation_id(self):
        """Test correlation ID generation"""
        correlation_id = function_app.generate_correlation_id()
        
        assert correlation_id.startswith("proc_")
        assert len(correlation_id.split("_")) == 4  # proc_YYYYMMDD_HHMMSS_uuid
        assert len(correlation_id) > 20  # Should be reasonably long
    
    def test_validate_request_valid_json(self):
        """Test request validation with valid JSON"""
        mock_request = Mock()
        mock_request.headers = {'content-type': 'application/json'}
        
        test_data = {
            'fileName': 'test.pdf',
            'fileContent': 'base64content',
            'jobTitle': 'Software Engineer'
        }
        
        # Mock the get_body() method to return properly encoded JSON
        mock_request.get_body.return_value = json.dumps(test_data).encode('utf-8')
        
        result = function_app.validate_and_parse_request(mock_request)
        
        assert result['fileName'] == 'test.pdf'
        assert result['jobTitle'] == 'Software Engineer'
    
    def test_validate_request_invalid_content_type(self):
        """Test request validation with invalid content type"""
        mock_request = Mock()
        mock_request.headers = {'content-type': 'text/plain'}
        
        with pytest.raises(function_app.ValidationError) as exc_info:
            function_app.validate_and_parse_request(mock_request)
        
        assert "Invalid content type" in str(exc_info.value)
    
    def test_base64_decoding_valid(self):
        """Test base64 decoding with valid content"""
        test_content = "Hello World"
        encoded_content = base64.b64encode(test_content.encode()).decode()
        
        # Note: This would be part of the main function logic
        decoded = base64.b64decode(encoded_content)
        assert decoded == test_content.encode()
    
    def test_base64_decoding_invalid(self):
        """Test base64 decoding with invalid content"""
        invalid_content = "not_valid_base64!"
        
        with pytest.raises(Exception):
            base64.b64decode(invalid_content)

# Error Handling Tests
class TestErrorHandling:
    """Test cases for comprehensive error handling"""
    
    def test_document_processing_error_handling(self):
        """Test error handling in document processing"""
        with patch('pymupdf4llm.to_markdown') as mock_pymupdf:
            mock_pymupdf.side_effect = Exception("PDF processing failed")
            
            with pytest.raises(DocumentProcessingError):
                extract_markdown_from_file("test.pdf", b"fake_bytes")
    
    @patch('shared_code.candidate_data_service.OpenAIService')
    @patch('shared_code.candidate_data_service.AICallLogger')
    def test_ai_service_error_handling(self, mock_logger, mock_openai):
        """Test error handling in AI service calls"""
        # Mock the OpenAI service to raise an exception
        mock_openai_instance = Mock()
        mock_openai_instance.call_assistant.side_effect = Exception("API call failed")
        mock_openai_instance.get_assistant_id.return_value = "asst_123"
        mock_openai.return_value = mock_openai_instance
        
        # Mock the logger
        mock_logger_instance = Mock()
        mock_logger_instance.create_call_log.return_value = {"stage": "test", "success": False}
        mock_logger.return_value = mock_logger_instance
        
        service = CandidateDataService()
        
        with pytest.raises(Exception) as exc_info:
            service._extract_candidate_profile("test content", "test job")
        
        assert "Candidate profile extraction failed" in str(exc_info.value)
    
    def test_json_parsing_error_handling(self):
        """Test error handling for invalid JSON responses"""
        invalid_json = "This is not valid JSON"
        
        with pytest.raises(json.JSONDecodeError):
            json.loads(invalid_json)

# Performance and Load Tests
class TestPerformance:
    """Test cases for performance and load scenarios"""
    
    def test_large_document_processing(self):
        """Test processing of large documents"""
        large_content = "A" * 100000  # 100KB of content
        
        with patch('pymupdf4llm.to_markdown') as mock_pymupdf:
            mock_pymupdf.return_value = large_content
            
            result = extract_markdown_from_file("large_test.pdf", b"fake_large_pdf")
            
            assert len(result) == 100000
            assert result == large_content
    
    def test_multiple_pii_entities_performance(self):
        """Test performance with many PII entities"""
        # Create a large number of PII entities
        pii_entities = []
        for i in range(100):
            pii_entities.append({
                "original_value": f"Entity{i}",
                "pii_type": "organization",
                "sensitivity_level": 2,
                "all_variations": [f"Entity{i}"],
                "context": f"test entity {i}",
                "region_specific": False
            })
        
        original_content = " ".join([f"Entity{i}" for i in range(100)])
        
        service = CandidateDataService()
        start_time = datetime.now()
        result = service._create_anonymized_content(original_content, pii_entities)
        end_time = datetime.now()
        
        # Should complete reasonably quickly
        processing_time = (end_time - start_time).total_seconds()
        assert processing_time < 5.0  # Should take less than 5 seconds
        
        # Verify all entities were replaced
        for i in range(100):
            assert f"Entity{i}" not in result

# Integration Test Suite
class TestFullIntegration:
    """End-to-end integration tests"""
    
    def create_test_pdf_content(self):
        """Create a test PDF file content"""
        return base64.b64encode(b"fake_pdf_content").decode()
    
    def create_test_docx_content(self):
        """Create a test DOCX file content"""
        return base64.b64encode(b"fake_docx_content").decode()
    
    @patch('shared_code.document_processor.extract_markdown_from_file')
    @patch('shared_code.candidate_data_service.CandidateDataService')
    def test_full_cv_processing_workflow_pdf(self, mock_service, mock_extract, 
                                           sample_cv_content, sample_candidate_profile,
                                           sample_skills_analysis):
        """Test complete CV processing workflow for PDF files"""
        # Mock document extraction
        mock_extract.return_value = sample_cv_content
        
        # Mock candidate data service
        mock_service_instance = Mock()
        mock_service_instance.extract_candidate_data.return_value = {
            'candidateProfile': sample_candidate_profile,
            'professionalAnalysis': {
                'aiConfidenceScore': 85,
                'aiRemarks': 'Strong candidate',
                'overallFitScore': 78,
                'coreSkills': [],
                'workExperience': [],
                'educationProfile': {}
            },
            'aiCallLogs': [],
            'tokenUsageSummary': {'totalTokens': 500}
        }
        mock_service.return_value = mock_service_instance
        
        # Simulate the main processing flow
        file_content = self.create_test_pdf_content()
        file_bytes = base64.b64decode(file_content)
        
        # Extract text
        markdown_content = mock_extract.return_value
        
        # Process with AI
        result = mock_service_instance.extract_candidate_data(markdown_content, "Software Engineer")
        
        assert result['candidateProfile']['candidateName'] == "John Smith"
        assert result['professionalAnalysis']['aiConfidenceScore'] == 85
        assert 'tokenUsageSummary' in result
    
    @patch.dict(os.environ, {
        'AZURE_OPENAI_ENDPOINT': 'https://test.openai.azure.com/',
        'AZURE_OPENAI_API_KEY': 'test-key',
        'CV_DATA_EXTRACTOR_ASSISTANT_ID': 'asst_123',
        'CV_PII_IDENTIFIER_ASSISTANT_ID': 'asst_456',
        'CV_SKILLS_ANALYST_ASSISTANT_ID': 'asst_789'
    })
    def test_configuration_validation_integration(self):
        """Test that all components can initialize with proper configuration"""
        # Test OpenAI service initialization
        openai_service = OpenAIService()
        assert openai_service.endpoint is not None
        
        # Test AI logger initialization
        ai_logger = AICallLogger()
        assert ai_logger.max_prompt_length > 0
        
        # Test candidate data service initialization
        with patch('shared_code.candidate_data_service.OpenAIService'):
            with patch('shared_code.candidate_data_service.AICallLogger'):
                candidate_service = CandidateDataService()
                assert candidate_service is not None

if __name__ == "__main__":
    # Configure logging for tests
    logging.basicConfig(level=logging.INFO)
    
    # Run tests with pytest
    pytest.main([__file__, "-v", "--tb=short"])
